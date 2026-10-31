"""XLSX export of a project: corpus + analysis results + human re-read + parameters."""

import asyncio
import hashlib
import io
import json
import os
import re
import uuid
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse, Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app import crate, prisma, references
from app.api.deps import get_current_user, get_db
from app.api.v1.projects import (
    STATUS_PATTERN,
    _require_lead_or_admin,
    _require_project_member,
    article_filter,
)
from app.core.config import get_settings
from app.db.models import (
    Article,
    ArticleFile,
    ArticlePrediction,
    ExportJob,
    HumanAnnotation,
    ImportJob,
    InferenceJob,
    ModelRegistry,
    PromptTemplate,
    Search,
    User,
)
from app.schema import project_fields
from app.searches import ScopedProject, default_search, get_search, in_search, lot_condition
from app.sources import source_label, source_name

router = APIRouter(prefix="/projects", tags=["export"])

DbDep = Annotated[AsyncSession, Depends(get_db)]
UserDep = Annotated[User, Depends(get_current_user)]

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_dir() -> str:
    """Directory of the generated zips (backend/worker shared volume)."""
    return os.environ.get("EXPORT_DIR", "/app/exports")


_MAX_CELL = 32000  # limite Excel : 32 767 caractères par cellule


def _cell(v):
    """Excel-compatible value: lists/dicts flattened, text cleaned and truncated."""
    if v is None or isinstance(v, bool | int | float):
        return v
    if isinstance(v, list):
        v = "\n".join(str(x) for x in v)
    elif isinstance(v, dict):
        v = "; ".join(f"{k}: {x}" for k, x in v.items())
    return crate.neutralize(ILLEGAL_CHARACTERS_RE.sub("", str(v))[:_MAX_CELL])


def _sheet(ws, header: list[str], rows: list[list], widths: dict[int, int] | None = None) -> None:
    ws.append(header)
    for r in rows:
        ws.append([_cell(v) for v in r])
    fill = PatternFill("solid", fgColor="DDE7F5")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), fill
        c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 18)


def _model_rows(label: str, name: str | None, snap: dict | None) -> list[list]:
    """ "Project" lines of a stage: name, identifier, version and snapshot details."""
    snap = snap or {}
    rows = [
        [label, name],
        [f"{label} · identifiant", snap.get("identifier") or "non enregistré"],
        [f"{label} · version", snap.get("version") or "non enregistrée"],
    ]
    keys = ("endpoint", "parameter_size", "quantization", "family", "format", "size", "modified_at")
    details = {k: snap[k] for k in keys if snap.get(k)}
    if details:
        rows.append([f"{label} · détails", details])
    return rows


def _import_source(job) -> str:
    """Database queried by an import; references / crate: mixed sources (Source column)."""
    if job.source:
        return source_name(job.source)
    return "sources par article" if job.kind in ("references", "crate") else "non enregistrée"


def build_workbook(project, articles, preds, annos, import_jobs, inf_job, user_email) -> bytes:
    fields = project_fields(project)
    wb = Workbook()

    # --- Results: one line per article -----------------------------------
    header = ["PMID", "PMCID", "DOI", "Titre", "Date de publication", "Journal"]
    header += ["Premier auteur", "Dernier auteur", "Auteurs", "Types de publication", "Langue"]
    header += ["Lien DOI", "Lien texte intégral / PDF", "Source", "Identifiant source"]
    header += ["Texte analysé", "Statut IA", "Pertinence IA", "Modèle"]
    header += ["Identifiant du modèle", "Version du modèle", "Décision IA"]
    header += [f"IA · {f['label']}" for f in fields]
    header += [f"Certitude · {f['label']}" for f in fields]
    header += ["Extraits (IA)", "Raisonnement (IA)", "Annotations humaines", "Statut final"]
    header += ["Relu par", "Statut de relecture", "Relu le"]
    header += [f"Humain · {f['label']}" for f in fields] + ["Notes humaines", "Résumé (abstract)"]
    rows = []
    for a in articles:
        p, hs = preds.get(a.id), annos.get(a.id, [])
        pd_ = (p.decision_json or {}) if p else {}
        h = hs[-1] if hs else None
        hd = (h.decision_json or {}) if h else {}
        final = "ANNOTÉ" if h else (p.status.value if p and p.status else "NON TRIÉ")
        au = a.authors or []
        row = [a.pmid, a.pmcid, a.doi, a.title, a.pub_date, a.journal]
        row += [au[0] if au else None, au[-1] if len(au) > 1 else None, "; ".join(au)]
        row += ["; ".join(a.pub_types or []), a.original_language]
        row += [f"https://doi.org/{a.doi}" if a.doi else None, a.fulltext_url]
        row += [source_label(a.source, a.native_id), a.native_id]
        row += [a.fulltext_source or ("abstract" if a.abstract else "—")]
        row += [p.status.value if p and p.status else None, p.confidence_score if p else None]
        row += [p.model_used if p else None]
        row += [p.model_identifier if p else None, p.model_version if p else None]
        row += [pd_.get("decision")]
        row += [pd_.get(f["key"]) for f in fields]
        row += [(pd_.get("field_confidence") or {}).get(f["key"]) for f in fields]
        row += [(p.evidence_quotes if p else None), pd_.get("reasoning_summary")]
        row += [len(hs), final]
        row += [
            getattr(h, "reviewer", None),
            hd.get("review_status"),
            f"{h.created_at:%Y-%m-%d %H:%M}" if h and h.created_at else None,
        ]
        row += [hd.get(f["key"]) for f in fields]
        row += [h.notes if h else None, a.abstract]
        rows.append(row)
    ws = wb.active
    ws.title = "Résultats"
    _sheet(ws, header, rows, {4: 60, 5: 12, 6: 28, 9: 40})
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (4, 9))
            # clickable links (DOI, full text); never javascript:/file: etc.
            if c.column in (12, 13) and str(c.value or "").startswith(("http://", "https://")):
                c.hyperlink = c.value
                c.style = "Hyperlink"

    # --- Summary ------------------------------------------------------------------
    counts: dict[str, int] = {}
    for row in rows:
        counts[row[header.index("Statut final")]] = (
            counts.get(row[header.index("Statut final")], 0) + 1
        )
    ws2 = wb.create_sheet("Résumé")
    _sheet(ws2, ["Statut final", "Articles"], [[k, v] for k, v in sorted(counts.items())], {1: 28})
    ws2.append([])
    ws2.append(["Total", len(rows)])
    ws2.append(["Avec texte intégral", sum(1 for a in articles if a.sections)])

    # --- Project / method (reproducibility) -----------------------------------
    ws3 = wb.create_sheet("Projet")
    info = [
        ["Projet", project.title],
        ["Question", project.question],
        ["Parties du texte lues", ", ".join(project.text_sections or ["methods"])],
        ["Fenêtres de lecture", project.read_windows],
        ["Exporté le", datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")],
        ["Exporté par", user_email],
        [],
        ["Champs de réponse", "Type / options"],
    ]
    info += [
        [f["label"], f"{f['type']} {', '.join(f.get('options') or [])}".strip()] for f in fields
    ]
    info += [[], ["Imports", "Date · source · articles importés"]]
    info += [
        [
            j.query,
            f"{j.created_at:%Y-%m-%d} · {_import_source(j)} · "
            f"{(j.result_json or {}).get('imported', '')}",
        ]
        for j in import_jobs
    ]
    if inf_job is not None:
        res = inf_job.result_json or {}
        info += [[], ["Dernière analyse", inf_job.created_at.strftime("%Y-%m-%d %H:%M")]]
        snaps = inf_job.model_info or {}  # instantané figé ; jamais le registre vivant
        info += _model_rows("Modèle (étage 2)", inf_job.model_used, snaps.get("main"))
        if inf_job.prefilter_model:
            info += _model_rows(
                "Pré-filtre (étage 1)", inf_job.prefilter_model, snaps.get("prefilter")
            )
        else:
            info.append(["Pré-filtre (étage 1)", inf_job.prefilter_model])
        info += [[k, v] for k, v in res.items()]
    for r in info:
        ws3.append([_cell(v) for v in r])
    ws3.column_dimensions["A"].width, ws3.column_dimensions["B"].width = 34, 90
    for c in ws3["A"]:
        c.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _load(db: AsyncSession, project, where=None, search=None):
    """Articles + last prediction / annotations per article + jobs, those of ``search`` (default:
    the whole initial search)."""
    search = search or await default_search(db, project)
    articles = (
        (
            await db.execute(
                select(Article)
                .where(*(where or [lot_condition(search)]))
                .order_by(Article.created_at, Article.id)
            )
        )
        .scalars()
        .all()
    )
    ids = [a.id for a in articles]
    preds: dict = {}
    annos: dict = {}
    if ids:
        for p in (
            await db.execute(
                select(ArticlePrediction)
                .where(
                    ArticlePrediction.article_id.in_(ids),
                    in_search(ArticlePrediction.search_id, search),
                )
                .order_by(ArticlePrediction.created_at.nulls_first(), ArticlePrediction.id)
            )
        ).scalars():
            preds[p.article_id] = p  # le plus récent écrase
        if search.kind == "blind":
            preds = {}  # échantillon aveugle : les réponses de l'IA ne sortent pas
        for h in (
            await db.execute(
                select(HumanAnnotation)
                .where(
                    HumanAnnotation.article_id.in_(ids),
                    in_search(HumanAnnotation.search_id, search),
                )
                .order_by(HumanAnnotation.created_at.nulls_first(), HumanAnnotation.id)
            )
        ).scalars():
            annos.setdefault(h.article_id, []).append(h)
        reviewers = {
            u.id: u
            for u in (
                await db.execute(
                    select(User).where(
                        User.id.in_({h.user_id for hs in annos.values() for h in hs if h.user_id})
                    )
                )
            ).scalars()
        }
        salt = get_settings().anon_salt
        for hs in annos.values():
            for h in hs:
                u = reviewers.get(h.user_id)
                h.reviewer = (u.full_name or u.email) if u else None  # name for results.xlsx
                # metadata.json: project-scoped pseudonym, derived from the same salted
                # digest kept at account deletion (``annotator_hash``).
                digest = h.annotator_hash or (
                    hashlib.sha256((salt + u.email).encode("utf-8")).hexdigest() if u else None
                )
                h.reviewer_ref = crate.reviewer_reference(project.id, digest)
    import_jobs = (
        (
            await db.execute(
                select(ImportJob)
                .where(
                    ImportJob.project_id == project.id,
                    or_(ImportJob.kind.is_(None), ImportJob.kind != "fulltext"),
                )
                .order_by(ImportJob.created_at)
            )
        )
        .scalars()
        .all()
    )
    inf_job = (
        (
            await db.execute(
                select(InferenceJob)
                .where(
                    InferenceJob.project_id == project.id,
                    in_search(InferenceJob.search_id, search),
                )
                .order_by(InferenceJob.created_at.desc())
                .limit(1)
            )
        )
        .scalars()
        .first()
    )
    return articles, preds, annos, import_jobs, inf_job


def _slug(project) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", project.title).strip("-")[:40] or "projet"


@router.get("/{project_id}/export.xlsx")
async def export_xlsx(
    project_id: uuid.UUID, db: DbDep, user: UserDep, search: uuid.UUID | None = None
) -> Response:
    project = await _require_project_member(db, project_id, user)
    scope = await get_search(db, project, search)
    articles, preds, annos, import_jobs, inf_job = await _load(db, project, search=scope)
    project = ScopedProject(project, scope)
    data = await asyncio.to_thread(
        build_workbook, project, articles, preds, annos, import_jobs, inf_job, user.email
    )
    return Response(
        data,
        media_type=_XLSX,
        headers={"Content-Disposition": f'attachment; filename="diggypaper-{_slug(project)}.xlsx"'},
    )


def _notes(articles, preds: dict, fields: list[dict]) -> dict:
    """Note per article: AI status, relevance and answers (to recover the sort in Zotero)."""
    out = {}
    for a in articles:
        p = preds.get(a.id)
        if p is None:
            continue
        dec = p.decision_json or {}
        answers = "; ".join(f"{f['label']}: {dec[f['key']]}" for f in fields if f["key"] in dec)
        rel = f" ({round(p.confidence_score * 100)} %)" if p.confidence_score is not None else ""
        ident = f" ({p.model_identifier})" if p.model_identifier else ""
        version = p.model_version or "non enregistrée"
        out[a.id] = (
            f"Diggypaper : {p.status.value if p.status else ''}{rel}. {answers}"
            f" Modèle : {p.model_used}{ident}, version {version}."
        ).strip()
    return out


async def _references(
    project_id, db, user, status, answers, fmt: str, search=None, q=None, fulltext=None
) -> Response:
    project = await _require_project_member(db, project_id, user)
    scope = await get_search(db, project, search)
    where = article_filter(scope, status, answers, q, fulltext)[0]
    articles, preds, *_ = await _load(db, project, where, scope)
    project = ScopedProject(project, scope)
    notes = _notes(articles, preds, project_fields(project))
    if fmt == "ris":
        body, mime = references.to_ris(articles, notes), references.RIS_MIME
    else:
        body, mime = references.to_bibtex(articles, notes), references.BIBTEX_MIME
    return Response(
        body.encode("utf-8"),
        media_type=f"{mime}; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{_slug(project)}.{fmt}"'},
    )


async def _prisma_origin(db, project, scope) -> dict | None:
    """A search other than the initial one: its starting point and its criteria."""
    if scope.is_default:
        return None
    parent = await db.get(Search, scope.parent_id) if scope.parent_id else None
    lot = parent or scope
    size = (
        await db.execute(select(func.count()).select_from(Article).where(lot_condition(lot)))
    ).scalar_one()
    spec = dict(scope.parent_spec or {})
    if "ids" in spec:  # pas la liste des identifiants dans le diagramme / project.json
        spec["n_ids"] = len(spec.pop("ids"))
    return {
        "name": (parent or scope).name,
        "size": size,
        "criteria": spec,
        "labels": {f["key"]: f["label"] for f in (lot.fields or [])},
    }


PRISMA_LANGS = "^(fr|en|de|es)$"


@router.get("/{project_id}/prisma")
async def get_prisma(
    project_id: uuid.UUID, db: DbDep, user: UserDep, search: uuid.UUID | None = None
) -> dict:
    """Counts of the PRISMA 2020 flow diagram of a search."""
    project = await _require_project_member(db, project_id, user)
    scope = await get_search(db, project, search)
    articles, preds, annos, import_jobs, _ = await _load(db, project, search=scope)
    return prisma.prisma_counts(
        articles, preds, annos, import_jobs, await _prisma_origin(db, project, scope)
    )


@router.get("/{project_id}/prisma.svg")
async def get_prisma_svg(
    project_id: uuid.UUID,
    db: DbDep,
    user: UserDep,
    lang: Annotated[str | None, Query(pattern=PRISMA_LANGS)] = None,
    search: uuid.UUID | None = None,
) -> Response:
    """PRISMA 2020 diagram in SVG (language: ``lang`` if given, else the project's)."""
    project = await _require_project_member(db, project_id, user)
    scope = await get_search(db, project, search)
    articles, preds, annos, import_jobs, _ = await _load(db, project, search=scope)
    counts = prisma.prisma_counts(
        articles, preds, annos, import_jobs, await _prisma_origin(db, project, scope)
    )
    parts = prisma.identified_parts(articles, import_jobs)
    svg = prisma.render_svg(counts, lang or project.default_language or "fr", *parts)
    return Response(
        svg,
        media_type="image/svg+xml",
        headers={"Content-Disposition": f'inline; filename="{_slug(project)}-prisma.svg"'},
    )


@router.get("/{project_id}/prisma.pdf")
async def get_prisma_pdf(
    project_id: uuid.UUID,
    db: DbDep,
    user: UserDep,
    lang: Annotated[str | None, Query(pattern=PRISMA_LANGS)] = None,
    search: uuid.UUID | None = None,
) -> Response:
    """PRISMA 2020 diagram as a vectorial PDF (same content as the SVG)."""
    project = await _require_project_member(db, project_id, user)
    scope = await get_search(db, project, search)
    articles, preds, annos, import_jobs, _ = await _load(db, project, search=scope)
    counts = prisma.prisma_counts(
        articles, preds, annos, import_jobs, await _prisma_origin(db, project, scope)
    )
    parts = prisma.identified_parts(articles, import_jobs)
    pdf = prisma.svg_to_pdf(
        prisma.render_svg(counts, lang or project.default_language or "fr", *parts)
    )
    return Response(
        pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{_slug(project)}-prisma.pdf"'},
    )


@router.get("/{project_id}/export.ris")
async def export_ris(
    project_id: uuid.UUID,
    db: DbDep,
    user: UserDep,
    status: Annotated[str | None, Query(pattern=STATUS_PATTERN)] = None,
    answers: Annotated[str | None, Query(max_length=4000)] = None,
    search: uuid.UUID | None = None,
    q: Annotated[str | None, Query(max_length=200)] = None,
    fulltext: bool | None = None,
) -> Response:
    """RIS references (Zotero, EndNote, Mendeley...); same filters as the article list."""
    return await _references(project_id, db, user, status, answers, "ris", search, q, fulltext)


@router.get("/{project_id}/export.bib")
async def export_bibtex(
    project_id: uuid.UUID,
    db: DbDep,
    user: UserDep,
    status: Annotated[str | None, Query(pattern=STATUS_PATTERN)] = None,
    answers: Annotated[str | None, Query(max_length=4000)] = None,
    search: uuid.UUID | None = None,
    q: Annotated[str | None, Query(max_length=200)] = None,
    fulltext: bool | None = None,
) -> Response:
    """BibTeX references; same filters as the article list."""
    return await _references(project_id, db, user, status, answers, "bib", search, q, fulltext)


async def write_crate(
    db: AsyncSession, project, user_email: str, tmp, progress=None, search=None
) -> dict:
    """Writes the RO-Crate (zip) into the binary file ``tmp`` -> {articles, pdf, text}.

    ``progress(done, total)`` (async, optional) is called every 10 articles.
    """
    search = search or await default_search(db, project)
    articles, preds, annos, import_jobs, inf_job = await _load(db, project, search=search)
    fields = project_fields(project)
    # Living registry: only as a fallback, for an old job without a snapshot.
    names = {n for n in (inf_job.model_used, inf_job.prefilter_model) if n} if inf_job else set()
    models = {}
    if names and not inf_job.model_info:
        rows = await db.execute(select(ModelRegistry).where(ModelRegistry.name.in_(names)))
        models = {m.name: m for m in rows.scalars()}
    ptid = inf_job.prompt_template_id if inf_job else None
    prompt = await db.get(PromptTemplate, ptid) if ptid else None
    info = crate.project_info(project, fields, import_jobs, inf_job, models, prompt)
    keys = crate.assign_keys(articles)
    xlsx = await asyncio.to_thread(
        build_workbook, project, articles, preds, annos, import_jobs, inf_job, user_email
    )

    files: dict[str, list[str]] = {}
    n_pdf = n_text = 0
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:

        def put(path: str, data: bytes | str, stored: bool = False) -> None:
            z.writestr(
                path,
                data,
                compress_type=zipfile.ZIP_STORED if stored else zipfile.ZIP_DEFLATED,
            )

        for i, a in enumerate(articles, 1):
            key, folder = keys[a.id], f"articles/{keys[a.id]}/"
            rec = crate.article_record(a, key, preds.get(a.id), annos.get(a.id, []), fields)
            present = ["metadata.json"]
            put(folder + "metadata.json", json.dumps(rec, ensure_ascii=False, indent=2))
            if a.sections:
                n_text += 1
                present.append("fulltext.md")
                put(folder + "fulltext.md", crate.fulltext_markdown(a))
            elif a.abstract:
                present.append("abstract.md")
                put(folder + "abstract.md", crate.fulltext_markdown(a))
            pdf = (
                await db.execute(select(ArticleFile.data).where(ArticleFile.article_id == a.id))
            ).scalar_one_or_none()
            if pdf:
                n_pdf += 1
                present.append("fulltext.pdf")
                put(folder + "fulltext.pdf", pdf, stored=True)
            files[key] = present
            if progress and i % 10 == 0:
                await progress(i, len(articles))
        now = datetime.now(UTC).isoformat(timespec="seconds")
        counts = prisma.prisma_counts(
            articles, preds, annos, import_jobs, await _prisma_origin(db, project, search)
        )
        info["prisma"] = counts
        put(
            "prisma.svg",
            prisma.render_svg(
                counts,
                project.default_language or "fr",
                *prisma.identified_parts(articles, import_jobs),
            ),
        )
        put("README.md", crate.readme(project, len(articles), n_pdf, n_text))
        put("project.json", json.dumps(info, ensure_ascii=False, indent=2))
        notes = _notes(articles, preds, fields)
        put("references.ris", references.to_ris(articles, notes))
        put("references.bib", references.to_bibtex(articles, notes))
        put("extraction.csv", crate.extraction_csv(articles, keys, preds, annos, fields))
        put("results.xlsx", xlsx)
        graph = crate.build_graph(project, articles, keys, files, info, now)
        put("ro-crate-metadata.json", json.dumps(graph, ensure_ascii=False, indent=2))
    return {"articles": len(articles), "pdf": n_pdf, "text": n_text}


def _export_payload(job: ExportJob) -> dict:
    return {
        "id": str(job.id),
        "status": job.status,
        "result": job.result_json,
        "error": job.error,
        "created_at": job.created_at.isoformat() if job.created_at else None,
        "download": job.status == "COMPLETED" and bool(job.filename),
    }


async def _latest_export(db: AsyncSession, project_id: uuid.UUID, search=None) -> ExportJob | None:
    """Last export of the project (or of this search if ``search`` is given)."""
    cond = [ExportJob.project_id == project_id]
    if search is not None:
        cond.append(in_search(ExportJob.search_id, search))
    return (
        (
            await db.execute(
                select(ExportJob)
                .where(*cond)
                .order_by(ExportJob.created_at.desc(), ExportJob.id.desc())
                .limit(1)
            )
        )
        .scalars()
        .first()
    )


@router.post("/{project_id}/export/crate", status_code=202)
async def start_crate_export(
    project_id: uuid.UUID, db: DbDep, user: UserDep, search: uuid.UUID | None = None
) -> dict:
    """Launches the RO-Crate export in the background (e-mail at the end); only one at a time."""
    from app.tasks.export import export_task

    project = await _require_project_member(db, project_id, user)
    await _require_lead_or_admin(db, project_id, user)
    job = await _latest_export(db, project.id)
    scope = await get_search(db, project, search)
    if (
        job is not None
        and job.status in ("PENDING", "RUNNING")
        and (job.search_id or (await default_search(db, project)).id) != scope.id
    ):  # un export à la fois par projet : ne pas renvoyer celui d'une autre recherche sans le dire
        raise HTTPException(status_code=409, detail="another export is in progress")
    if job is None or job.status not in ("PENDING", "RUNNING"):
        job = ExportJob(
            project_id=project.id, user_id=user.id, search_id=scope.id, status="PENDING"
        )
        db.add(job)
        await db.commit()
        await db.refresh(job)
        export_task.delay(str(job.id))
        await db.refresh(job)
    return _export_payload(job)


@router.get("/{project_id}/export/crate")
async def get_crate_export(
    project_id: uuid.UUID, db: DbDep, user: UserDep, search: uuid.UUID | None = None
) -> dict | None:
    """Last RO-Crate export of the search (status, progress, downloadable?)."""
    project = await _require_project_member(db, project_id, user)
    job = await _latest_export(db, project_id, await get_search(db, project, search))
    return _export_payload(job) if job else None


@router.get("/{project_id}/export/crate/{job_id}/download")
async def download_crate(
    project_id: uuid.UUID, job_id: uuid.UUID, db: DbDep, user: UserDep
) -> FileResponse:
    project = await _require_project_member(db, project_id, user)
    job = await db.get(ExportJob, job_id)
    if job is None or job.project_id != project.id or job.status != "COMPLETED" or not job.filename:
        raise HTTPException(status_code=404, detail="export not found")
    path = Path(export_dir()) / job.filename
    if not path.is_file():  # purgé (plus de 7 jours) : à relancer
        raise HTTPException(status_code=410, detail="export expired")
    return FileResponse(
        path, media_type="application/zip", filename=f"{_slug(project)}-ro-crate.zip"
    )
