"""Invitation d'équipe par e-mail + export XLSX du projet."""

import asyncio
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.security import decode_invite_token
from app.db.models import (
    Article,
    ArticlePrediction,
    PredictionStatus,
    Project,
    ProjectMember,
    ProjectMemberRole,
)


@pytest.fixture
async def lead_project(session, researcher_user, client, do_login):
    await do_login("researcher@lab.fr", "researcher-pass-123", consent=True)
    project = Project(title="Revue test", created_by=researcher_user.id, question="Q ?")
    session.add(project)
    await session.commit()
    session.add(
        ProjectMember(
            project_id=project.id,
            user_id=researcher_user.id,
            role_in_project=ProjectMemberRole.LEAD,
        )
    )
    await session.commit()
    return project


@pytest.fixture
def with_mailer(monkeypatch, mock_mailer):
    async def fake(db):
        return mock_mailer

    monkeypatch.setattr("app.api.v1.team.mailer_from_db", fake)
    return mock_mailer


async def test_invite_unknown_email_sends_register_link_then_joins(
    client, lead_project, with_mailer
):
    resp = await client.post(
        f"/api/v1/projects/{lead_project.id}/members/invite",
        json={"email": "new@lab.fr", "role": "ANNOTATOR", "language": "fr"},
    )
    assert resp.status_code == 201
    assert resp.json() == {"status": "invited", "email_sent": True, "link": None}
    body = with_mailer.sent[-1]["body"]
    token = body.split("token=")[1].split()[0]
    claims = decode_invite_token(token)
    assert claims["project_id"] == str(lead_project.id) and claims["role"] == "ANNOTATOR"
    client.cookies.clear()
    reg = await client.post(
        "/api/v1/auth/register", json={"token": token, "password": "new-pass-1234"}
    )
    assert reg.status_code == 201
    await client.post(
        "/api/v1/auth/login",
        json={"email": "new@lab.fr", "password": "new-pass-1234", "consent": True},
    )
    members = (await client.get(f"/api/v1/projects/{lead_project.id}/members")).json()["items"]
    assert {m["email"]: m["role"] for m in members}["new@lab.fr"] == "ANNOTATOR"


async def test_invite_existing_user_added_directly(client, session, lead_project, with_mailer):
    from app.db.models import UserRole
    from tests.conftest import _make_user

    await _make_user(session, "old@lab.fr", UserRole.RESEARCHER, "old-pass-1234")
    resp = await client.post(
        f"/api/v1/projects/{lead_project.id}/members/invite", json={"email": "OLD@lab.fr"}
    )
    assert resp.json()["status"] == "invited"  # réponse uniforme (pas d'énumération)
    again = await client.post(
        f"/api/v1/projects/{lead_project.id}/members/invite", json={"email": "old@lab.fr"}
    )
    assert again.status_code == 409


async def test_invite_without_smtp_returns_link(client, lead_project, monkeypatch):
    async def none(db):
        return None

    monkeypatch.setattr("app.api.v1.team.mailer_from_db", none)
    resp = await client.post(
        f"/api/v1/projects/{lead_project.id}/members/invite", json={"email": "x@lab.fr"}
    )
    assert resp.json()["email_sent"] is False and "/register?token=" in resp.json()["link"]


async def test_cannot_remove_last_lead(client, lead_project, researcher_user):
    resp = await client.delete(f"/api/v1/projects/{lead_project.id}/members/{researcher_user.id}")
    assert resp.status_code == 409


async def test_export_xlsx_contains_results(client, session, lead_project):
    art = Article(
        project_id=lead_project.id,
        pmid="42",
        title="Mice",
        abstract="abs",
        doi="10/x",
        journal="Nature",
        pub_date="2025-03",
        authors=["Doe J", "Roe R", "Poe P"],
        fulltext_url="https://repo.org/a.pdf",
    )
    session.add(art)
    await session.commit()
    session.add(
        ArticlePrediction(
            article_id=art.id,
            model_used="coder",
            decision_json={
                "use_ai_behavior": True,
                "reasoning_summary": "because",
                "field_confidence": {"use_ai_behavior": 0.81},
            },
            confidence_score=0.93,
            evidence_quotes=["quote one", "quote two"],
            status=PredictionStatus.AUTO_ACCEPTED,
        )
    )
    await session.commit()
    resp = await client.get(f"/api/v1/projects/{lead_project.id}/export.xlsx")
    assert resp.status_code == 200 and "spreadsheetml" in resp.headers["content-type"]
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Résultats", "Résumé", "Projet"]
    rows = list(wb["Résultats"].values)
    head, first = rows[0], rows[1]
    assert first[head.index("PMID")] == "42"
    assert first[head.index("Statut IA")] == "AUTO_ACCEPTED"
    assert first[head.index("Pertinence IA")] == 0.93
    assert "quote one\nquote two" == first[head.index("Extraits (IA)")]
    assert first[head.index("Raisonnement (IA)")] == "because"
    assert first[head.index("Certitude · Uses AI behavior analysis")] == 0.81
    assert first[head.index("Premier auteur")] == "Doe J"
    assert first[head.index("Dernier auteur")] == "Poe P"
    assert first[head.index("Journal")] == "Nature"
    assert (await session.execute(select(Article))).scalars().first() is not None
    assert first[head.index("Lien DOI")] == "https://doi.org/10/x"
    assert first[head.index("Lien texte intégral / PDF")] == "https://repo.org/a.pdf"


async def test_list_articles_filters_by_answers(client, session, lead_project):
    import json

    lead_project.fields = [
        {"key": "ai", "label": "AI", "type": "boolean"},
        {"key": "sex", "label": "Sex", "type": "choice", "options": ["MALE", "FEMALE"]},
        {"key": "note", "label": "Note", "type": "text"},
    ]
    session.add(lead_project)
    arts = []
    for pmid, dec in (
        ("1", {"ai": True, "sex": "MALE", "note": "Deep_Lab Cut"}),
        ("2", {"ai": False, "sex": "FEMALE", "note": "manual"}),
    ):
        a = Article(project_id=lead_project.id, pmid=pmid, doi=f"10/{pmid}")
        session.add(a)
        await session.flush()
        session.add(
            ArticlePrediction(
                article_id=a.id,
                model_used="m",
                decision_json=dec,
                confidence_score=0.9,
                status=PredictionStatus.AUTO_ACCEPTED,
            )
        )
        arts.append(a)
    await session.commit()

    async def pmids(**flt):
        r = await client.get(
            f"/api/v1/projects/{lead_project.id}/articles", params={"answers": json.dumps(flt)}
        )
        assert r.status_code == 200, r.text
        return sorted(i["pmid"] for i in r.json()["items"])

    assert await pmids(ai=True) == ["1"]
    assert await pmids(ai=False) == ["2"]
    assert await pmids(sex="FEMALE") == ["2"]
    assert await pmids(note="deep_lab") == ["1"]
    assert await pmids(note="deep%") == []  # % littéral, pas un joker
    assert await pmids(ai=True, sex="FEMALE") == []
    assert await pmids(ai="") == ["1", "2"]
    bad = await client.get(
        f"/api/v1/projects/{lead_project.id}/articles", params={"answers": '{"zzz": 1}'}
    )
    assert bad.status_code == 422
    item = (await client.get(f"/api/v1/projects/{lead_project.id}/articles")).json()["items"][0]
    assert item["doi_url"].startswith("https://doi.org/10/")


async def test_export_ro_crate_async(
    client, engine, session, lead_project, mock_mailer, monkeypatch, tmp_path
):
    import json
    import zipfile

    from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

    from app.db.models import ArticleFile, InferenceJob, ModelRegistry
    from app.tasks import export as texport
    from app.tasks import inference as tinf

    a1 = Article(
        project_id=lead_project.id,
        pmid="42",
        title="Mice",
        abstract="abs",
        doi="10.1/x",
        journal="Nature",
        pub_date="2025-03",
        authors=["Doe J", "Roe R"],
        sections=[{"title": "Methods", "kind": "methods", "text": "mice were housed"}],
        fulltext_source="pdf:upload",
    )
    a2 = Article(project_id=lead_project.id, pmid="43", title="Rats", abstract="only abstract")
    session.add_all([a1, a2, ModelRegistry(name="Laya", model_identifier="http://l#ckpt")])
    await session.commit()
    session.add_all(
        [
            ArticleFile(article_id=a1.id, data=b"%PDF-1.4 fake"),
            ArticlePrediction(
                article_id=a1.id,
                model_used="Laya",
                decision_json={
                    "use_ai_behavior": True,
                    "field_confidence": {"use_ai_behavior": 0.9},
                },
                confidence_score=0.95,
                status=PredictionStatus.AUTO_ACCEPTED,
            ),
            InferenceJob(project_id=lead_project.id, model_used="Laya", result_json={"total": 2}),
        ]
    )
    await session.commit()

    queued = []
    monkeypatch.setattr(texport.export_task, "delay", lambda jid: queued.append(jid))
    monkeypatch.setenv("EXPORT_DIR", str(tmp_path))

    async def fake_mailer(_session):
        return mock_mailer

    monkeypatch.setattr(tinf, "_make_mailer", fake_mailer)
    base = f"/api/v1/projects/{lead_project.id}/export/crate"
    assert (await client.get(base)).json() is None
    r = await client.post(base)
    assert r.status_code == 202 and r.json()["status"] == "PENDING" and len(queued) == 1
    assert (await client.post(base)).json()["id"] == r.json()["id"]  # un seul à la fois
    assert len(queued) == 1
    assert (await client.get(f"{base}/{r.json()['id']}/download")).status_code == 404

    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    stats = await texport._run_export(queued[0], session_factory=factory)
    assert stats["articles"] == 2 and stats["pdf"] == 1 and stats["size"] > 0
    assert mock_mailer.sent and "RO-Crate" in mock_mailer.sent[0]["body"]

    status = (await client.get(base)).json()
    assert status["status"] == "COMPLETED" and status["download"] is True
    resp = await client.get(f"{base}/{status['id']}/download")
    assert resp.status_code == 200 and resp.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(resp.content))
    names = set(z.namelist())
    assert {"ro-crate-metadata.json", "project.json", "extraction.csv", "results.xlsx"} <= names
    assert "articles/pmid-42/fulltext.pdf" in names and "articles/pmid-42/fulltext.md" in names
    assert "articles/pmid-43/abstract.md" in names and "articles/pmid-43/fulltext.pdf" not in names
    assert z.read("articles/pmid-42/fulltext.pdf") == b"%PDF-1.4 fake"
    rec = json.loads(z.read("articles/pmid-42/metadata.json"))
    assert rec["ai"]["certainty"]["use_ai_behavior"] == 0.9 and rec["last_author"] == "Roe R"
    graph = json.loads(z.read("ro-crate-metadata.json"))["@graph"]
    ids = {e["@id"] for e in graph}
    assert {"./", "https://doi.org/10.1/x", "articles/pmid-42/fulltext.pdf", "#analysis"} <= ids
    root = next(e for e in graph if e["@id"] == "./")
    assert all(p["@id"] in ids for p in root["hasPart"])  # pas de référence orpheline
    assert root["license"]["@id"] == "https://creativecommons.org/licenses/by/4.0/"
    assert "https://creativecommons.org/licenses/by/4.0/" in ids  # licence décrite dans le graphe
    info = json.loads(z.read("project.json"))
    assert info["analysis"]["stage_2_model"]["config_sha256"]
    assert z.read("extraction.csv").decode("utf-8-sig").splitlines()[0].startswith("key,pmid")


async def test_crate_download_expired_410(client, session, lead_project, monkeypatch, tmp_path):
    from app.db.models import ExportJob

    monkeypatch.setenv("EXPORT_DIR", str(tmp_path))
    job = ExportJob(project_id=lead_project.id, status="COMPLETED", filename="gone.zip")
    session.add(job)
    await session.commit()
    r = await client.get(f"/api/v1/projects/{lead_project.id}/export/crate/{job.id}/download")
    assert r.status_code == 410


async def test_export_references_ris_bibtex_follow_filters(client, session, lead_project):
    arts = []
    for pmid, status, title in (
        ("1", PredictionStatus.AUTO_ACCEPTED, "Deep learning & mice_100%"),
        ("2", PredictionStatus.AUTO_REJECTED, "Off topic"),
    ):
        a = Article(
            project_id=lead_project.id,
            pmid=pmid,
            title=title,
            abstract="abs",
            doi=f"10.1/{pmid}",
            journal="Nature",
            pub_date="2025-03-12",
            authors=["van der Berg JA", "Roe R"],
        )
        session.add(a)
        await session.flush()
        session.add(
            ArticlePrediction(article_id=a.id, model_used="m", confidence_score=0.9, status=status)
        )
        arts.append(a)
    await session.commit()
    base = f"/api/v1/projects/{lead_project.id}"

    ris = await client.get(f"{base}/export.ris")
    assert ris.status_code == 200 and "research-info-systems" in ris.headers["content-type"]
    assert ris.text.count("TY  - JOUR") == 2 and ris.text.count("ER  - ") == 2
    assert "AU  - van der Berg, JA" in ris.text and "DA  - 2025/03/12" in ris.text
    assert "DO  - 10.1/1" in ris.text and "AN  - PMID:1" in ris.text

    only = await client.get(f"{base}/export.ris", params={"status": "AUTO_ACCEPTED"})
    assert only.text.count("TY  - JOUR") == 1 and "Off topic" not in only.text

    bib = await client.get(f"{base}/export.bib", params={"status": "AUTO_ACCEPTED"})
    assert "@article{pmid1," in bib.text and "pmid2" not in bib.text
    assert r"Deep learning \& mice\_100\%" in bib.text
    assert "author = {van der Berg, JA and Roe, R}" in bib.text and "year = {2025}" in bib.text
    assert (await client.get(f"{base}/export.ris", params={"status": "X"})).status_code == 422


async def test_prisma_counts_and_svg(client, session, lead_project):
    from app.db.models import HumanAnnotation, ImportJob, ImportJobStatus

    def art(pmid, sections=None):
        a = Article(project_id=lead_project.id, pmid=pmid, title=pmid, sections=sections)
        session.add(a)
        return a

    acc, rej, pend, human_no, _untriaged = (
        art("1", [{"title": "m", "kind": "methods", "text": "x"}]),
        art("2"),
        art("3"),
        art("4"),
        art("5"),
    )
    await session.flush()
    for a, st in (
        (acc, PredictionStatus.AUTO_ACCEPTED),
        (rej, PredictionStatus.AUTO_REJECTED),
        (pend, PredictionStatus.NEEDS_HUMAN_REVIEW),
        (human_no, PredictionStatus.NEEDS_HUMAN_REVIEW),
    ):
        session.add(ArticlePrediction(article_id=a.id, model_used="m", status=st))
    session.add(HumanAnnotation(article_id=human_no.id, decision_json={"confidence_score": 0.0}))
    session.add(
        ImportJob(
            project_id=lead_project.id,
            query="q",
            status=ImportJobStatus.COMPLETED,
            result_json={"imported": 5, "failed": 2},
        )
    )
    await session.commit()
    base = f"/api/v1/projects/{lead_project.id}"
    c = (await client.get(f"{base}/prisma")).json()
    assert c == {
        "identified": 7,
        "failed": 2,
        "duplicates": 0,
        "by_criteria": 0,
        "via_citations": 0,
        "origin": None,
        "not_screened": 1,
        "screened": 4,
        "excluded": 2,
        "excluded_ai": 1,
        "excluded_human": 1,
        "sought": 2,
        "not_retrieved": 1,
        "assessed": 1,
        "pending": 1,
        "included": 1,
    }
    svg = await client.get(f"{base}/prisma.svg", params={"lang": "en"})
    assert svg.headers["content-type"].startswith("image/svg+xml")
    assert "Records identified (PubMed, imported lists)" in svg.text and "n = 7" in svg.text
    pdf = await client.get(f"{base}/prisma.pdf", params={"lang": "fr"})
    assert pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
    from pypdf import PdfReader

    text = PdfReader(io.BytesIO(pdf.content)).pages[0].extract_text()
    assert "Références identifiées" in text and "n = 7" in text  # accents conservés

    # Recherche dérivée : le point de départ est la recherche d'origine, les critères sont
    # indiqués et les articles écartés ne sont pas des « doublons ».
    r = await client.post(
        f"{base}/searches", json={"name": "Retenus", "selection": {"status": "AUTO_ACCEPTED"}}
    )
    assert r.status_code == 201, r.text
    d = (await client.get(f"{base}/prisma", params={"search": r.json()["id"]})).json()
    assert (d["identified"], d["duplicates"], d["failed"], d["by_criteria"]) == (5, 0, 0, 4)
    assert d["origin"]["criteria"] == {"status": "AUTO_ACCEPTED"}
    svg = await client.get(f"{base}/prisma.svg", params={"search": r.json()["id"], "lang": "fr"})
    assert "Hors critères de sélection (n = 4)" in svg.text and "statut IA : accepté" in svg.text
    assert "Doublons" not in svg.text


async def test_resume_inference_skips_done_articles(engine, session):
    from datetime import UTC, datetime, timedelta

    from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

    from app.db.models import (
        InferenceJob,
        InferenceJobStatus,
        ModelRegistry,
        Project,
        ProjectStatus,
        PromptTemplate,
    )
    from app.tasks.inference import _run_inference

    project = Project(title="P", status=ProjectStatus.FROZEN)
    session.add(project)
    await session.commit()
    prompt = PromptTemplate(project_id=project.id, version_number="v1", system_prompt="S")
    done = Article(project_id=project.id, pmid="1", title="done", abstract="a")
    todo = Article(project_id=project.id, pmid="2", title="todo", abstract="a")
    session.add_all([prompt, done, todo, ModelRegistry(name="Kev", model_identifier="kev:q4")])
    await session.commit()
    before = datetime.now(UTC) - timedelta(minutes=5)
    session.add(
        ArticlePrediction(
            article_id=done.id,
            prompt_template_id=prompt.id,
            model_used="Kev",
            confidence_score=0.9,
            status=PredictionStatus.AUTO_ACCEPTED,
            created_at=datetime.now(UTC),
        )
    )
    job = InferenceJob(
        project_id=project.id,
        prompt_template_id=prompt.id,
        model_used="Kev",
        status=InferenceJobStatus.PENDING,
        resume_after=before,
    )
    session.add(job)
    await session.commit()

    class Chat:
        def __init__(self):
            self.calls = 0

        async def infer(self, model, system_prompt, user_prompt):
            self.calls += 1
            return {"confidence_score": 0.05, "evidence_quotes": [], "reasoning_summary": ""}

    chat = Chat()
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    stats = await _run_inference(str(job.id), session_factory=factory, inference=chat)
    assert chat.calls == 2  # 1 warm-up + seul l'article restant est analysé
    assert stats["processed"] == 2 and stats["auto_accepted"] == 1 and stats["auto_rejected"] == 1


async def test_trigger_inference_resume(client, session, lead_project, monkeypatch):
    from datetime import UTC, datetime

    from sqlalchemy import select

    from app.db.models import (
        InferenceJob,
        InferenceJobStatus,
        ModelRegistry,
        ModelShare,
        ProjectStatus,
    )
    from app.tasks import inference as tinf

    monkeypatch.setattr(
        tinf.inference_task, "delay", lambda jid: type("R", (), {"state": "PENDING"})()
    )
    lead_project.status = ProjectStatus.FROZEN
    kev = ModelRegistry(name="Kev", model_identifier="kev:q4")
    session.add_all([lead_project, kev])
    await session.commit()
    session.add(ModelShare(model_id=kev.id, project_id=lead_project.id))
    await session.commit()
    url = f"/api/v1/projects/{lead_project.id}/inference"
    body = {"model": "Kev", "resume": True}
    assert (await client.post(url, json=body)).status_code == 409  # rien à reprendre

    first = await client.post(url, json={"model": "Kev"})
    assert first.status_code == 202
    job = (await session.execute(select(InferenceJob))).scalars().one()
    job.status, job.created_at = InferenceJobStatus.FAILED, datetime(2026, 1, 1, tzinfo=UTC)
    await session.commit()

    again = await client.post(url, json=body)
    assert again.status_code == 202
    jobs = (await session.execute(select(InferenceJob).order_by(InferenceJob.created_at))).scalars()
    new = [j for j in jobs if str(j.id) == again.json()["job_id"]][0]
    assert new.resume_after is not None and new.resume_after.year == 2026
    assert new.prompt_template_id == job.prompt_template_id


async def test_review_status_reviewer_and_prefill(client, session, lead_project):
    lead_project.fields = [{"key": "ai", "label": "AI", "type": "boolean"}]
    session.add(lead_project)
    a = Article(project_id=lead_project.id, pmid="1")
    session.add(a)
    await session.flush()
    session.add(
        ArticlePrediction(
            article_id=a.id,
            model_used="m",
            decision_json={"ai": True},
            confidence_score=0.4,
            status=PredictionStatus.NEEDS_HUMAN_REVIEW,
        )
    )
    await session.commit()
    base = f"/api/v1/projects/{lead_project.id}/articles"
    assert (await client.get(f"{base}/{a.id}/annotation")).json() is None
    assert (await client.get(base)).json()["items"][0]["review"] is None

    r = await client.post(f"{base}/{a.id}/annotation", json={"ai": True})
    assert r.status_code == 201, r.text
    assert r.json()["status"] == "VALIDATED"
    item = (await client.get(base)).json()["items"][0]
    assert item["annotated"] and item["review"]["status"] == "VALIDATED"
    assert item["review"]["by"]  # nom ou e-mail du relecteur

    # correction : la dernière relecture l'emporte, préremplie au rechargement
    await asyncio.sleep(1.1)  # SQLite : created_at à la seconde, sinon égalité départagée au hasard
    r = await client.post(f"{base}/{a.id}/annotation", json={"ai": False})
    assert r.json()["status"] == "CORRECTED"
    got = (await client.get(f"{base}/{a.id}/annotation")).json()
    assert got["status"] == "CORRECTED" and got["decision_json"]["ai"] is False
    assert (await client.get(base)).json()["items"][0]["review"]["status"] == "CORRECTED"
    # la liste affiche les réponses de la relecture (pas celles du modèle) et filtre dessus
    item = (await client.get(base)).json()["items"][0]
    assert item["answers"] == [
        {"key": "ai", "label": "AI", "value": False, "confidence": None, "source": "human"}
    ]
    assert len((await client.get(base, params={"answers": '{"ai": false}'})).json()["items"]) == 1
    assert (await client.get(base, params={"answers": '{"ai": true}'})).json()["items"] == []

    x = await client.get(f"/api/v1/projects/{lead_project.id}/export.xlsx")
    ws = load_workbook(io.BytesIO(x.content))["Résultats"]
    head = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    assert row[head.index("Statut de relecture")] == "CORRECTED"
    assert row[head.index("Relu par")]
