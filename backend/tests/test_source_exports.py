"""Phase 2 multi-sources : la source de chaque article / import dans les exports (XLSX, RIS,
BibTeX, RO-Crate) et à l'import d'un crate (spec 2026-10-29-sources-phase2-identite §4)."""

import csv
import io
import json
import uuid
import zipfile

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.api.v1.export import write_crate
from app.db.models import Article, ImportJob, ImportJobStatus, Project, ProjectStatus, User
from app.tasks.crate_import import _import_crate

GRAPH = json.dumps({"@graph": [{"conformsTo": {"@id": "https://w3id.org/ro/crate/1.1"}}]})


async def _project(client, do_login) -> str:
    await do_login("researcher@lab.fr", "researcher-pass-123", consent=True)
    resp = await client.post("/api/v1/projects", json={"title": "Sources"})
    assert resp.status_code == 201
    return resp.json()["id"]


def _articles(pid) -> list[Article]:
    pid = uuid.UUID(str(pid))
    return [
        Article(
            project_id=pid,
            pmid="42",
            doi="10.1/a",
            source="pubmed",
            native_id="42",
            title="From PubMed",
        ),
        Article(
            project_id=pid,
            doi="10.1/y",
            source="crossref",
            native_id="10.1/y",
            title="From Crossref",
        ),
        Article(project_id=pid, doi="10.1/old", title="Legacy"),
    ]


async def test_xlsx_results_source_columns(client, session, researcher_user, do_login):
    """Résultats : « Source » (libellé lisible) et « Identifiant source » juste après le lien
    texte intégral ; vides pour un article ancien sans source."""
    pid = await _project(client, do_login)
    session.add_all(_articles(pid))
    await session.commit()
    resp = await client.get(f"/api/v1/projects/{pid}/export.xlsx")
    rows = list(load_workbook(io.BytesIO(resp.content))["Résultats"].values)
    head = rows[0]
    i = head.index("Source")
    assert i == head.index("Lien texte intégral / PDF") + 1
    assert head[i + 1] == "Identifiant source"
    by_title = {r[head.index("Titre")]: r[i : i + 2] for r in rows[1:]}
    assert by_title["From PubMed"] == ("PubMed PMID 42", "42")
    assert by_title["From Crossref"] == ("Crossref DOI 10.1/y", "10.1/y")
    assert by_title["Legacy"] == (None, None)


async def test_xlsx_projet_sheet_import_sources(client, session, researcher_user, do_login):
    """Projet : chaque import mentionne sa source ; un ancien job de recherche sans source dit
    « non enregistrée » ; un import de références (sources mêlées) renvoie au détail par article."""
    pid = await _project(client, do_login)
    session.add_all(
        [
            ImportJob(
                project_id=uuid.UUID(pid),
                query="cancer AND mice",
                source="pubmed",
                status=ImportJobStatus.COMPLETED,
                result_json={"imported": 12},
            ),
            ImportJob(
                project_id=uuid.UUID(pid),
                query="old query",
                status=ImportJobStatus.COMPLETED,
                result_json={"imported": 3},
            ),
            ImportJob(
                project_id=uuid.UUID(pid),
                query="references (2)",
                kind="references",
                status=ImportJobStatus.COMPLETED,
                result_json={"imported": 2},
            ),
        ]
    )
    await session.commit()
    resp = await client.get(f"/api/v1/projects/{pid}/export.xlsx")
    info = {r[0]: r[1] for r in load_workbook(io.BytesIO(resp.content))["Projet"].values if r}
    assert "Requêtes PubMed" not in info
    assert info["Imports"] == "Date · source · articles importés"
    assert "PubMed" in info["cancer AND mice"] and info["cancer AND mice"].endswith("· 12")
    assert "non enregistrée" in info["old query"]
    assert "sources par article" in info["references (2)"]


async def test_ris_db_tag_and_bibtex_source_field(client, session, researcher_user, do_login):
    """RIS : tag DB = nom de la source ; BibTeX : champ ``source`` lisible ; rien sans source."""
    pid = await _project(client, do_login)
    session.add_all(_articles(pid))
    await session.commit()
    ris = (await client.get(f"/api/v1/projects/{pid}/export.ris")).text
    db_lines = sorted(line for line in ris.splitlines() if line.startswith("DB  - "))
    assert db_lines == ["DB  - Crossref", "DB  - PubMed"]
    bib = (await client.get(f"/api/v1/projects/{pid}/export.bib")).text
    assert "source = {PubMed PMID 42}" in bib
    assert "source = {Crossref DOI 10.1/y}" in bib
    assert bib.count("source = {") == 2


async def test_crate_exports_source_and_job_source(session, tmp_path):
    """RO-Crate : metadata.json et extraction.csv portent source + native_id ; project.json
    donne la source de chaque import ; le graphe nomme l'import par sa source."""
    project = Project(title="Crate", question="Q ?")
    session.add(project)
    await session.flush()
    session.add_all(_articles(project.id))
    session.add(
        ImportJob(
            project_id=project.id,
            query="cancer",
            source="pubmed",
            status=ImportJobStatus.COMPLETED,
            result_json={"imported": 1},
        )
    )
    await session.commit()
    out = tmp_path / "c.zip"
    with open(out, "wb") as fh:
        await write_crate(session, project, "lead@lab.fr", fh)
    with zipfile.ZipFile(out) as z:
        rec = json.loads(z.read("articles/pmid-42/metadata.json"))
        rows = list(csv.reader(io.StringIO(z.read("extraction.csv").decode("utf-8-sig"))))
        info = json.loads(z.read("project.json"))
        graph = json.loads(z.read("ro-crate-metadata.json"))["@graph"]
    assert (rec["source"], rec["native_id"]) == ("pubmed", "42")
    head = rows[0]
    assert head[:2] == ["key", "pmid"]
    i = head.index("source")
    assert i == head.index("doi") + 1 and head[i + 1] == "native_id"
    by_title = {r[head.index("title")]: (r[i], r[i + 1]) for r in rows[1:]}
    assert by_title == {
        "From PubMed": ("pubmed", "42"),
        "From Crossref": ("crossref", "10.1/y"),
        "Legacy": ("", ""),
    }
    assert info["pubmed_queries"][0]["source"] == "pubmed"
    actions = [e for e in graph if e.get("@type") == "SearchAction"]
    assert [a["name"] for a in actions] == ["Import PubMed"]


async def test_crate_import_restores_source_and_defaults_old_crates(engine, session, tmp_path):
    """Import d'un crate : source + native_id restaurés ; crate ancien sans source → pubmed/PMID
    si PMID, sinon rien."""
    metas = {
        "doi-10-1-y": {
            "doi": "10.1/y",
            "title": "New",
            "source": "crossref",
            "native_id": "10.1/y",
        },
        "pmid-7": {"pmid": "7", "title": "Old with PMID"},
        "doi-10-1-z": {"doi": "10.1/z", "title": "Old without PMID"},
    }
    path = tmp_path / "old.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("ro-crate-metadata.json", GRAPH)
        z.writestr("project.json", json.dumps({"title": "Restored"}))
        for key, meta in metas.items():
            z.writestr(f"articles/{key}/metadata.json", json.dumps(meta | {"human": []}))
    new = Project(title="Restored", status=ProjectStatus.PROCESSING)
    importer = User(email="me@lab.fr", hashed_password="x")
    session.add_all([new, importer])
    await session.commit()
    new_id = new.id
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    result = await _import_crate(str(new_id), str(path), str(importer.id), session_factory=factory)
    assert result["imported"] == 3
    session.expire_all()
    arts = (await session.execute(select(Article).where(Article.project_id == new_id))).scalars()
    got = {a.title: (a.source, a.native_id) for a in arts}
    assert got == {
        "New": ("crossref", "10.1/y"),
        "Old with PMID": ("pubmed", "7"),
        "Old without PMID": (None, None),
    }


def test_source_label_and_name():
    """Libellés (module pur ``app.sources``) : nom de la source et identifiant lisible."""
    from app.sources import source_label, source_name

    assert source_name("pubmed") == "PubMed" and source_name("crossref") == "Crossref"
    assert source_name("europepmc") == "europepmc"  # inconnu : code brut
    assert source_name(None) is None
    assert source_label("pubmed", "123") == "PubMed PMID 123"
    assert source_label("crossref", "10.1/x") == "Crossref DOI 10.1/x"
    assert source_label("file", None) == "Fichier importé"
    assert source_label("europepmc", "W1") == "europepmc W1"
    assert source_label(None, "123") is None
